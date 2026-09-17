#!/usr/bin/env python3
"""Build the 13-min collection-efficiency deck (16:9, English, speaker notes).

One line of logic per slide, the figure does the talking: the deck is built
around one figure, a single pulse read as ADC, as current and as power, and then
follows that pulse to the whole detector.

    python3 build_pptx.py            # writes SNOLAB_R4_collection_efficiency_20260910.pptx
    module load libreoffice && libreoffice --headless --convert-to pdf <the pptx>

Every number is copied from ../results/plots/current_power_overlay/*.txt or
../NOTES.md; nothing is computed here.
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from PIL import Image

HERE = os.path.dirname(os.path.abspath(__file__))
FIG = os.path.join(HERE, "figures")
OUT = os.path.join(HERE, "SNOLAB_R4_collection_efficiency_20260910.pptx")

NAVY = RGBColor(0x1F, 0x38, 0x64)
DARK = RGBColor(0x33, 0x33, 0x33)
GRAY = RGBColor(0x66, 0x66, 0x66)
RED = RGBColor(0xC0, 0x39, 0x2B)
SW, SH = Inches(13.333), Inches(7.5)

prs = Presentation()
prs.slide_width, prs.slide_height = SW, SH
BLANK = prs.slide_layouts[6]
n_slide = 0
L, T, W = Inches(0.55), Inches(1.25), Inches(12.3)


def textbox(s, l, t, w, h):
    tb = s.shapes.add_textbox(l, t, w, h)
    tb.text_frame.word_wrap = True
    return tb


def run(p, text, size, color=DARK, bold=False, italic=False):
    r = p.add_run(); r.text = text
    r.font.size, r.font.color.rgb, r.font.name = Pt(size), color, "Arial"
    r.font.bold, r.font.italic = bold, italic
    return r


def title(s, text, sub=None):
    tb = textbox(s, L, Inches(0.25), W, Inches(0.8))
    run(tb.text_frame.paragraphs[0], text, 27, NAVY, bold=True)
    if sub:
        # a (plain, bold) pair puts the second half of the subtitle in bold
        para = tb.text_frame.add_paragraph()
        if isinstance(sub, tuple):
            run(para, sub[0], 14, GRAY)
            run(para, sub[1], 14, DARK, bold=True)
        else:
            run(para, sub, 14, GRAY)
    ln = s.shapes.add_shape(1, L, Inches(1.06), W, Emu(18000))
    ln.fill.solid(); ln.fill.fore_color.rgb = NAVY; ln.line.fill.background()


def lines(s, items, l, t, w, h, size=18):
    """A few short lines. A line starting with '!' is the red take-away."""
    tb = textbox(s, l, t, w, h)
    tf = tb.text_frame
    for i, txt in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(8)
        red = txt.startswith("!")
        run(p, txt.lstrip("!"), size, RED if red else DARK, bold=red)


def pic(s, name, l, t, max_w, max_h):
    path = os.path.join(FIG, name)
    w0, h0 = Image.open(path).size
    scale = min(max_w / w0, max_h / h0)
    w, h = int(w0 * scale), int(h0 * scale)
    s.shapes.add_picture(path, l + int((max_w - w) / 2), t + int((max_h - h) / 2),
                         width=w, height=h)


def new(ttl, sub=None, notes=""):
    global n_slide
    n_slide += 1
    s = prs.slides.add_slide(BLANK)
    if ttl:
        title(s, ttl, sub)
    tb = textbox(s, Inches(12.5), Inches(7.05), Inches(0.6), Inches(0.35))
    run(tb.text_frame.paragraphs[0], str(n_slide), 10, GRAY)
    s.notes_slide.notes_text_frame.text = notes
    return s


# ------------------------------------------------------------------ 1 title
s = new(None, notes="Title.")
tb = textbox(s, Inches(0.9), Inches(2.3), Inches(11.5), Inches(1.8))
run(tb.text_frame.paragraphs[0], "Phonon collection efficiency of Z7", 40, NAVY, bold=True)
tb = textbox(s, Inches(0.9), Inches(4.7), Inches(11.5), Inches(1.2))
run(tb.text_frame.paragraphs[0], "Zhiheng Li  ·  University of Minnesota  ·  SuperCDMS SNOLAB Run 4  ·  September 2026", 16)
run(tb.text_frame.add_paragraph(), "github.com/ZhihengLi0/urop  →  snolab/collection_efficiency/", 13, GRAY)

# ------------------------------------------------------------------ 2 question
s = new("The question", notes="Efficiency = energy absorbed by the TESs, summed over all "
        "channels, / 10.37 keV, for a detector operated at 0 V (no NTL amplification). Z7, "
        "K-line events of the Ge activation, SNOLAB R4, 2207 events over 30 series.")
pic(s, "chain.png", L, Inches(1.4), W, Inches(2.6))
lines(s, [
    "!For detectors operated at 0 V (no NTL effect):",
    "Collection efficiency  =  energy absorbed by the TESs, summed over all channels  ÷  the 10.37 keV of the event",
    "Sample: Z7, the K-line events of the Ge activation run, from SNOLAB R4 (2207 events, 30 series, all 11 channels).",
], L, Inches(4.3), W, Inches(2.7))

# ------------------------------------------------------------------ 3 the formula
s = new("The formula, and the two numbers it needs", "Method 1 of the CDMS collection-efficiency note",
        notes="TES small-signal result (Irwin & Hilton), Method 1 of the CDMS note. Coefficients "
              "from the measured bias point of each channel: I0(R_L-R0) = 0.45 uV, 2R_L = 0.0385 Ohm "
              "for PBS1. For a two-exponential pulse the integral is a formula in A, tau_f, tau_r.")
pic(s, "formula.png", L, T, W, Inches(4.3))
lines(s, [
    "δP = the change of Joule heating when the current moves by δI: a linear term and a small quadratic term.",
    "The note also gives Method 2. In our sign convention it reads δP = I₀(R_L − R₀)·δI − R_L·(δI)²: the same linear term, only the quadratic term differs, and it gives 1.7 % less energy (PBS1, median of 15 events).",
], L, Inches(5.6), W, Inches(1.6), size=15)

# ------------------------------------------------------------------ 4 the figure
s = new("Example: one pulse",
        ("Z7 PBS1, event 30646:  ", "the same trace as ADC counts, as current, as power"),
        notes="Top: the pulse zoomed on its peak, three axes. Bottom: the power and its area. "
              "1 ADC = 0.318 nA. The next two slides walk through the two panels.")
pic(s, "peak_full.png", L, T, W, Inches(5.3))
lines(s, [
    "!The two-exponential is fitted to the 100 kHz low-passed trace (blue), not to the raw samples; the raw samples are only drawn.",
], L, Inches(6.62), W, Inches(0.5), size=14)

# ------------------------------------------------------------------ 5 height
s = new("The height: take it from the fit, not from the largest sample",
        notes="Grey raw samples, blue 100 kHz low pass (what the fit sees), green 20 kHz (official "
              "prefilter), red the two-exponential fit, pink band fit ± 1 sigma (89 ADC). Yellow: 98 "
              "samples within 10% of the peak. Raw max 885 ADC vs fit 717: +23%. The maximum of 98 "
              "equal samples is the biggest upward noise. The fit averages every sample.")
pic(s, "peak_top.png", L, T, W, Inches(4.95))
lines(s, [
    "Yellow band: 98 samples sit within 10 % of the peak. The largest of them is the largest noise excursion.",
    "!Largest raw sample 885 ADC, fit 717 ADC: the maximum reads 23 % high. Filtering helps; the fit removes it.",
], L, Inches(6.3), W, Inches(1.0), size=15)

# ------------------------------------------------------------------ 6 power & area
s = new("The power pulse, and its area is the energy",
        notes="Red: power from the fitted current, two terms added. Blue dashed: linear term alone. "
              "Purple dotted: quadratic term x20. Peak 104.6 fW = 102.6 + 2.0: quadratic 1.9% of "
              "peak, 1.3% of energy. Shaded area 305.2 eV, 99% inside this window. Bottom line: "
              "the bias point that fixes the coefficients.")
pic(s, "peak_bottom.png", L, T, W, Inches(4.0))
lines(s, [
    "Red: the fitted current put through the formula. Blue dashed: the linear term alone. Purple: the quadratic term, ×20.",
    "!The quadratic term is 1.9 % of the peak — the power pulse has the shape of the current pulse.",
    "!The shaded area is the energy this channel absorbed: 305 eV.",
], L, Inches(5.4), W, Inches(1.8), size=15)

# ------------------------------------------------------------------ 7 fit vs raw
s = new("Why the fitted pulse is integrated, and not the data",
        notes="Top: raw current of events 30646 and 210571, y zoomed to +-60 nA; dashed green is the "
              "mean after the pulse. Bottom: accumulated energy. Left: after-pulse mean +0.1 nA, raw "
              "316 vs fit 305 eV. Right: -6.2 nA, i.e. -2.8 fW over 24 ms = -420 eV, raw -169 vs fit 277.")
pic(s, "raw_cum_slide.png", L, T, W, Inches(5.1))
lines(s, [
    "Top: the raw current of the two events; the small panel zooms in on the baseline after the pulse. Bottom: the energy accumulated from the raw current (grey) and from the fit (red).",
    "!Right: after the pulse the current sits 6 nA below its baseline — hidden in 28 nA of noise, yet over 24 ms it is −420 eV: raw −169 eV, fit 277 eV.",
], L, Inches(6.33), W, Inches(0.9), size=13)

# ------------------------------------------------------------------ 8 all channels
s = new("The same event in its other channels: they add up to 33 %",
        "Z7 event 30646; four of the eleven channels",
        notes="PES1 489 eV, PES2 195 eV: position. Sum over 11 channels 3419 eV = 33.0%. Official "
              "window on the raw traces: 3397 eV = 32.8%. PDS2: slow swing after 5 ms, a "
              "low-frequency artefact; residual 0.14, three times the others.")
pic(s, "allchan_slide.png", L, T, W, Inches(4.65))
lines(s, [
    "The channels share the energy unevenly (PES1 489 eV, PES2 195 eV): the event sat closer to one side.",
    "!Summed over the eleven channels: 3419 eV = 33.0 % of 10.37 keV.   PDS2 (bottom right) carries a slow swing.",
], L, Inches(6.0), W, Inches(1.1), size=15)

# ------------------------------------------------------------------ 9 all events
s = new("Now every event: one channel gives a wide, lopsided distribution",
        "PBS1, 1917 K-line events, one fit and one integral each",
        notes="Red: fit. Grey outline: official window on the raw trace; same peak to 0.2%. Peak 285 "
              "eV, width 42 eV = 15%, for a fixed-energy line. Right tail to 600 eV: events near "
              "PBS1. The width is position, not noise.")
pic(s, "hist_PBS1.png", L, T, Inches(8.8), Inches(5.9))
lines(s, [
    "Red: from the fit. Grey: the official window on the raw trace. Same peak to 0.2 %.",
    "!Peak 285 eV, width 15 % — for a line whose energy is fixed.",
    "The tail to the right is the events that happened near this channel.",
    "!One channel's width is event position, not noise.",
], Inches(9.35), Inches(1.6), Inches(3.6), Inches(5.0), size=15)

# ------------------------------------------------------------------ 10 sum
s = new("Summed over the channels: 32 ± 1 %",
        "the total energy the TESs absorbed per K-line event",
        notes="Grey: ten channels that fit ≥ 95% of events, 1827 events, 3162 eV = 30.5%, width "
              "7.5% — the position dependence cancels in the sum. PDS2 fits only 42% (the swing) "
              "and on those events the other ten channels are 4.8% low, so its share is a bracket: "
              "31.5 to 32.9%. Red line 10.37 keV.")
pic(s, "hist_sum.png", L, T, Inches(8.8), Inches(5.9))
lines(s, [
    "!Ten channels summed, 1827 events: 3162 eV = 30.5 %, and the width drops to 7.5 % — the position effect cancels.",
    "PDS2 fits only 42 % of the events (that slow swing), and those events are not typical; its share can only be bracketed.",
    "!Full efficiency 31.5–32.9 %:  32 ± 1 %.",
    "Red line: 10.37 keV — two thirds of the energy never reaches the TESs.",
], Inches(9.35), Inches(1.6), Inches(3.6), Inches(5.0), size=15)

# ------------------------------------------------------------------ 11 summary
s = new("Summary, and what is still open",
        notes="Chain: current -> power -> energy -> sum -> 32 ± 1%. R37 CUTE note: 26-41%. Open: "
              "Z7 only; HV bias unknown (0 V assumed); no dIdV, inductor term dropped; PDS2; 12.5% "
              "of events fit nowhere.")
lines(s, [
    "A current pulse becomes a power pulse through the TES equations; the area under the power is the energy. The quadratic term is 2 %.",
    "The energy comes from the fitted pulse, integrated from −∞ to +∞: no bias from the noisy maximum, no drift.",
    "One channel: 15 % wide and lopsided (position). Ten channels summed: 7.5 % and symmetric.",
    "!Z7 absorbs 32 ± 1 % of a 10.37 keV event in its TESs. The CDMS note on the R37 CUTE tower reports 26–41 %.",
    "Open: Z7 only · the HV bias of Z7 is not established (0 V assumed) · no dI/dV, so the inductor term is dropped · PDS2 is the ± 1 % · 12.5 % of events fit in no channel and are not in the histograms.",
], L, Inches(1.45), W, Inches(5.5), size=17)

# ================================================================== backup
s = new("Backup — the events", notes="Z7 summed PTOFamps, our rebuild with no cut; K line at 2e-6 A.")
pic(s, "spectrum_zip7.png", L, T, W, Inches(5.75))

s = new("Backup — the three versions of the power formula",
        notes="c2 = 2 Method 1 (used), 1 exact, -1 Method 2. Relative to exact: +0.58%, -1.15%.")
lines(s, [
    "δP = I₀(R_L − R₀)·δI + c₂·R_L·(δI)²     with c₂ = 2 (Method 1, used here), 1 (exact small-signal), −1 (Method 2)",
    "Relative to the exact form: Method 1 +0.58 %, Method 2 −1.15 % in energy — far below the 15 % single-channel width and the PDS2 bracket.",
    "The inductor term of the full equations needs L and the loop gain from dI/dV, which do not exist for these series on MSI; it is neglected.",
    "The official Eabs (trigger bin 16383, baseline bins 93..15758, 5-pole 20 kHz prefilter, window −0.5/+1 ms, same formula) is reproduced to 0.2 %; the grey outlines on the histogram slides are it.",
], L, Inches(1.45), W, Inches(5.5), size=17)

s = new("Backup — how the formulas are derived")
pic(s, "derivation_backup_en.png", L, T, W, Inches(5.85))

for ttl, name, note in [
    ("Backup — 15 events, fitted current and power (PBS1)", "pulse_15events.png", "All 15 example events."),
    ("Backup — 15 events, cumulative energy (PBS1)", "cum_15events.png", "Fit vs raw, 15 events."),
    ("Backup — event 30646, all channels", "allchan_pulses.png", "Sum 3419 eV = 33.0%."),
    ("Backup — energy distribution of every channel", "hist_allchan.png",
     "Every channel is right-skewed; the peaks must not be added (21% low), the means add. PDS2 left out."),
]:
    s = new(ttl, notes=note)
    pic(s, name, L, T, W, Inches(5.75))

prs.save(OUT)
print(f"saved {OUT}: {n_slide} slides")
